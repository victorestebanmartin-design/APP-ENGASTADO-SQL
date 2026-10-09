import io
import json
import sqlite3
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.excel_manager import ExcelManager, _parse_instrucciones
from app.mangueras_editor import leer_preparacion, exportar_preparacion, serializar_instrucciones


def excel_original():
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Format'
    hoja.append(['Cable / Marca', 'Sección', 'Cod. cable', 'De Elemento', 'Observaciones', 'Longitud'])
    hoja.append(['MG1', '2x0,5', 'C1', 'X1', '<-PM100/MRS20 // PM200->', '=10+20'])
    hoja.append(['MG1', '12X1', 'C1', 'X1', 'Nota', 250])
    hoja.append(['Hilo', '0,5', 'C2', 'X2', None, 100])
    hoja.append(['MG3', '3 x 1', 'C3', 'X3', None, 100])
    hoja['A2'].font = Font(bold=True)
    libro.create_sheet('Otra').append(['No modificar'])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def test_sugerencias_y_legacy():
    datos = leer_preparacion(excel_original(), 'corte.xlsx')
    assert [fila['fila'] for fila in datos['filas']] == [2, 3, 5]
    assert datos['filas'][0]['de']['m_mrs_medida'] == 20
    assert datos['filas'][0]['para']['pm'] == 200


def test_descarga_compatible_sin_perder_excel(tmp_path):
    original = excel_original()
    datos = leer_preparacion(original, 'corte.xlsx')
    antes = deepcopy(datos)
    cambio = datos['filas'][1]
    cambio['de'] = _parse_instrucciones('PM120/MRC30/A100/A1_80/A2_90')
    cambio['para'] = _parse_instrucciones('PM150/M_CORTAR')
    cambio['retractil_de'] = [{'codigo': '649255', 'medida': 40}, {'codigo': '649251', 'medida': 70}]
    salida = exportar_preparacion(original, 'corte.xlsx', [cambio], datos['revision'])
    libro = load_workbook(salida)
    assert libro['Format']['F2'].value == '=10+20'
    assert libro['Format']['A2'].font.bold
    assert libro['Otra']['A1'].value == 'No modificar'
    assert libro['Format']['E3'].value == 'Nota'
    destino = tmp_path / 'editado.xlsx'
    libro.save(destino)
    mangueras = ExcelManager(str(tmp_path)).get_mangueras('editado.xlsx')
    assert len(mangueras) == 2
    assert mangueras[0]['de']['pm'] == 100
    assert mangueras[1]['de'] == cambio['de']
    assert mangueras[1]['retractil_de'] == cambio['retractil_de']
    assert leer_preparacion(original, 'corte.xlsx') == antes


def test_borrar_legacy_conserva_notas():
    original = excel_original()
    datos = leer_preparacion(original, 'corte.xlsx')
    cambio = datos['filas'][0]
    cambio['de'] = _parse_instrucciones('')
    cambio['para'] = _parse_instrucciones('')
    salida = exportar_preparacion(original, 'corte.xlsx', [cambio], datos['revision'])
    assert load_workbook(salida)['Format']['E2'].value is None


@pytest.mark.parametrize('token', ['PM0/M0/A0/A1_0', 'PM120/MRS/A25', 'PM120/MRC/A100', 'M_CORTAR'])
def test_roundtrip_todos_los_campos(token):
    instrucciones = _parse_instrucciones(token)
    assert _parse_instrucciones(serializar_instrucciones(instrucciones)) == instrucciones


def test_rechaza_revision_o_fila_incorrecta():
    original = excel_original()
    datos = leer_preparacion(original, 'corte.xlsx')
    with pytest.raises(ValueError, match='ha cambiado'):
        exportar_preparacion(original, 'corte.xlsx', [], 'vieja')
    with pytest.raises(ValueError, match='Fila'):
        exportar_preparacion(original, 'corte.xlsx', [{'fila': 4}], datos['revision'])


@pytest.mark.parametrize('valor', [-1, 1.5, True, 'abc', 1000000])
def test_rechaza_medidas_invalidas(valor):
    with pytest.raises(ValueError, match='medidas'):
        serializar_instrucciones({'pm': valor})


def test_editor_y_descarga_http(client):
    pagina = client.get('/mangueras/editor')
    assert pagina.status_code == 200
    assert b'\x00' not in pagina.data
    original = excel_original()
    respuesta = client.post('/api/mangueras/editor/leer', data={'excel': (io.BytesIO(original), 'corte.xlsx')})
    assert respuesta.status_code == 200
    datos = respuesta.get_json()
    cambio = datos['filas'][1]
    cambio['de']['pm'] = 80
    descarga = client.post('/api/mangueras/editor/descargar', data={
        'excel': (io.BytesIO(original), 'corte.xlsx'), 'revision': datos['revision'],
        'cambios': json.dumps([cambio]),
    })
    assert descarga.status_code == 200
    assert 'corte_preparacion.xlsx' in descarga.headers['Content-Disposition']
    libro = load_workbook(io.BytesIO(descarga.data))
    assert libro['Format']['G3'].value == 'PM80'


def test_corte_registrado_no_se_sobrescribe(client, app):
    ruta = Path(app.config['UPLOAD_FOLDER']) / 'corte.xlsx'
    original = excel_original()
    ruta.write_bytes(original)
    datos = client.post('/api/mangueras/editor/leer', data={'archivo': 'corte.xlsx'}).get_json()
    respuesta = client.post('/api/mangueras/editor/descargar', data={
        'archivo': 'corte.xlsx', 'revision': datos['revision'], 'cambios': '[]',
    })
    assert respuesta.status_code == 200
    assert ruta.read_bytes() == original
    ruta.write_bytes(original + b'cambio')
    respuesta = client.post('/api/mangueras/editor/descargar', data={
        'archivo': 'corte.xlsx', 'revision': datos['revision'], 'cambios': '[]',
    })
    assert respuesta.status_code == 400
    assert 'ha cambiado' in respuesta.get_json()['error']


@pytest.mark.parametrize('data', [{'archivo': '../fuera.xlsx'}, {'cambios': '{'}])
def test_editor_rechaza_origen_invalido(client, data):
    assert client.post('/api/mangueras/editor/descargar', data=data).status_code == 400


def test_conserva_tokens_y_retractiles_no_reconocidos():
    libro = load_workbook(io.BytesIO(excel_original()))
    hoja = libro['Format']
    hoja['G1'] = 'Instrucciones Mangueras DE'
    hoja['G3'] = 'PM100/OTRA'
    hoja['H1'] = 'Retráctil DE'
    hoja['H3'] = 'COD_ABC_20/pendiente'
    buffer = io.BytesIO()
    libro.save(buffer)
    original = buffer.getvalue()
    datos = leer_preparacion(original, 'corte.xlsx')
    cambio = datos['filas'][1]
    assert cambio['de']['otros_tokens'] == ['OTRA']
    cambio['de']['pm'] = 200
    salida = exportar_preparacion(original, 'corte.xlsx', [cambio], datos['revision'])
    editado = load_workbook(salida)['Format']
    assert editado['G3'].value == 'PM200/OTRA'
    assert editado['H3'].value == 'COD_ABC_20/pendiente'


def test_editor_respeta_permisos_mangueras(client, monkeypatch):
    import app.auth as auth
    from app.routes import base
    monkeypatch.setattr(auth, 'gate_operario_activo', lambda: True)
    monkeypatch.setattr(auth, '_pc_configurado_o_redirect', lambda: ('mangueras', None))
    monkeypatch.setattr(auth, '_operario_en_sesion_valido', lambda: 'Sin permiso')
    monkeypatch.setattr(base, 'operario_puede', lambda *args: False)
    assert client.get('/mangueras/editor').status_code == 403
    for accion in ('leer', 'descargar'):
        respuesta = client.post(f'/api/mangueras/editor/{accion}')
        assert respuesta.status_code == 403
        assert respuesta.get_json()['success'] is False


@pytest.mark.parametrize('instrucciones', [
    {'m_cortar': True, 'm_mrc': True}, {'m': 20, 'm_mrs': True},
    {'a_especificos': {'1': None}}, {'a_especificos': {'0': 20}},
])
def test_rechaza_preparacion_ambigua(instrucciones):
    with pytest.raises(ValueError):
        serializar_instrucciones(instrucciones)


def test_exportacion_mantiene_celdas_como_texto():
    original = excel_original()
    datos = leer_preparacion(original, 'corte.xlsx')
    cambio = datos['filas'][1]
    cambio['de']['otros_tokens'] = ['=SUM(A1)']
    salida = exportar_preparacion(original, 'corte.xlsx', [cambio], datos['revision'])
    celda = load_workbook(salida)['Format']['G3']
    assert celda.value == '=SUM(A1)'
    assert celda.data_type == 's'


def test_exportacion_normaliza_cabeceras_de_preparacion(tmp_path):
    libro = load_workbook(io.BytesIO(excel_original()))
    hoja = libro['Format']
    hoja['G1'] = 'instrucciones mangueras de'
    hoja['G3'] = 'PM10'
    buffer = io.BytesIO()
    libro.save(buffer)
    original = buffer.getvalue()
    datos = leer_preparacion(original, 'corte.xlsx')
    cambio = datos['filas'][1]
    cambio['de']['pm'] = 50
    salida = exportar_preparacion(original, 'corte.xlsx', [cambio], datos['revision'])
    destino = tmp_path / 'editado.xlsx'
    destino.write_bytes(salida.getvalue())
    assert ExcelManager(str(tmp_path)).get_mangueras('editado.xlsx')[1]['de']['pm'] == 50


def test_aplicar_actualiza_mismo_excel_y_conserva_backup(admin_client, app):
    ruta = Path(app.config['UPLOAD_FOLDER']) / 'corte.xlsx'
    original = excel_original()
    ruta.write_bytes(original)
    manager = ExcelManager(str(ruta.parent))
    assert manager.get_mangueras('corte.xlsx')[0]['de']['pm'] == 100
    datos = leer_preparacion(original, 'corte.xlsx')
    cambio = datos['filas'][0]
    cambio['de']['pm'] = 180
    with sqlite3.connect(app.config['DB_PATH']) as conexion:
        bd_antes = list(conexion.iterdump())
    respuesta = admin_client.post('/api/mangueras/editor/aplicar', data={
        'destino': 'corte.xlsx', 'revision': datos['revision'], 'cambios': json.dumps([cambio]),
    })
    assert respuesta.status_code == 200
    resultado = respuesta.get_json()
    assert resultado['archivo'] == 'corte.xlsx'
    assert resultado['datos']['revision'] != datos['revision']
    assert manager.get_mangueras('corte.xlsx')[0]['de']['pm'] == 180
    assert (ruta.parent / resultado['backup']).read_bytes() == original
    libro = load_workbook(ruta)
    assert libro['Format']['F2'].value == '=10+20'
    assert libro['Otra']['A1'].value == 'No modificar'
    with sqlite3.connect(app.config['DB_PATH']) as conexion:
        assert list(conexion.iterdump()) == bd_antes


def test_aplicar_exige_administracion(client, app):
    ruta = Path(app.config['UPLOAD_FOLDER']) / 'corte.xlsx'
    original = excel_original()
    ruta.write_bytes(original)
    assert client.post('/api/mangueras/editor/aplicar', data={'destino': 'corte.xlsx'}).status_code == 401
    assert ruta.read_bytes() == original


@pytest.mark.parametrize('tipo', ['revision', 'medida', 'destino', 'json'])
def test_aplicar_error_no_modifica_original(admin_client, app, tipo):
    ruta = Path(app.config['UPLOAD_FOLDER']) / 'corte.xlsx'
    original = excel_original()
    ruta.write_bytes(original)
    datos = leer_preparacion(original, 'corte.xlsx')
    cambio = datos['filas'][0]
    cambio['de']['pm'] = -1 if tipo == 'medida' else 180
    respuesta = admin_client.post('/api/mangueras/editor/aplicar', data={
        'destino': '../fuera.xlsx' if tipo == 'destino' else 'corte.xlsx',
        'revision': 'vieja' if tipo == 'revision' else datos['revision'],
        'cambios': '{' if tipo == 'json' else json.dumps([cambio]),
    })
    assert respuesta.status_code == 400
    assert ruta.read_bytes() == original
    assert list(ruta.parent.iterdir()) == [ruta]


def test_aplicar_fallo_de_escritura_no_modifica_original(admin_client, app, monkeypatch):
    from app.routes import manguitos
    ruta = Path(app.config['UPLOAD_FOLDER']) / 'corte.xlsx'
    original = excel_original()
    ruta.write_bytes(original)
    datos = leer_preparacion(original, 'corte.xlsx')
    def fallar(*args):
        raise OSError('Error de escritura simulado')
    monkeypatch.setattr(manguitos.os, 'replace', fallar)
    respuesta = admin_client.post('/api/mangueras/editor/aplicar', data={
        'destino': 'corte.xlsx', 'revision': datos['revision'], 'cambios': '[]',
    })
    assert respuesta.status_code == 500
    assert ruta.read_bytes() == original
    assert not any(fichero.name.startswith('tmp') for fichero in ruta.parent.iterdir())