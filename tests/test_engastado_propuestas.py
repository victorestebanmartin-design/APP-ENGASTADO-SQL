"""Propuestas de engastado: el operario avisa de que un terminal no se pone en
un cable, el admin lo aprueba, se escribe el '*' y se hereda en la ED nueva."""
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app.engastado_bloqueos import aplicar_bloqueos

CABECERA = ['Cable / Marca', 'Sección', 'Cod. cable', 'De Elemento', 'De Terminal',
            'Para Elemento', 'Para Terminal', 'Longitud']
FILAS = [
    ['W1', '0,5', 'C1', 'X1', 'T1', 'Y1', 'T2', 100],
    ['W2', '0,5', 'C1', 'X1', 'T1', 'Y2', 'T1', 100],
    ['W3', '0,5', 'C2', 'X2', 'T3', 'Y3', 'T1', 100],
]


def excel(filas, extra_arriba=0):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Format'
    hoja.append(CABECERA)
    for i in range(extra_arriba):
        hoja.append([f'N{i}', '0,5', 'C9', 'Z9', 'T9', 'Z8', 'T8', 50])
    for fila in filas:
        hoja.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def preparar(app, admin_client):
    from app.routes.base import db
    with db.engine.begin() as conn:
        conn.execute(text("INSERT INTO operarios (id, nombre) VALUES ('autor', 'Operario')"))
        conn.execute(text("INSERT INTO operarios (id, nombre, modulos_permitidos) "
                          "VALUES ('jefe', 'Jefe', '[\"admin\",\"engastado\"]')"))
        conn.execute(text(
            "INSERT INTO operario_logins (id, operario_nombre, modulo, activo, timestamp_login, ultimo_latido) "
            "VALUES ('login-autor', 'Operario', 'engastado', 1, :ahora, :ahora)"
        ), {'ahora': datetime.now().isoformat()})
    assert admin_client.put('/api/operarios/jefe/credencial-admin', json={'clave': 'personal123'}).status_code == 200
    ruta = Path(app.config['UPLOAD_FOLDER']) / 'corte.xlsx'
    ruta.write_bytes(excel(FILAS))
    operario = app.test_client()
    with operario.session_transaction() as sesion:
        sesion['operario_actual'] = 'Operario'
        sesion['operario_login_id'] = 'login-autor'
    jefe = app.test_client()
    assert jefe.post('/admin/login', data={'operario_id': 'jefe', 'clave': 'personal123'}).status_code == 302
    return ruta, operario, jefe


def avisar(operario, cables, terminal='T1'):
    return operario.post('/api/engastado/propuestas', json={
        'archivo': 'corte.xlsx', 'terminal': terminal, 'paquete': 'X1 (C1)', 'cables': cables,
        'motivo_categoria': 'Plano / documentacion', 'motivo': '',
    })


def cables_de_t1(client):
    datos = client.get('/api/datos_trabajo_v3?archivo=corte.xlsx&terminal=T1').get_json()
    return {(d['cable'], lado) for p in datos['paquetes'] for lado, lista in p['detalles_cables'].items()
            for d in lista}


def test_aviso_exige_operario_identificado(client):
    assert avisar(client, [{'fila': 2, 'lado': 'de'}]).status_code == 401


def test_detalles_de_engastado_llevan_la_fila_del_excel(app, admin_client):
    preparar(app, admin_client)
    datos = admin_client.get('/api/datos_trabajo_v3?archivo=corte.xlsx&terminal=T1').get_json()
    filas = {d['cable']: d['fila'] for p in datos['paquetes'] for lista in p['detalles_cables'].values() for d in lista}
    assert filas == {'W1': 2, 'W2': 3, 'W3': 4}


def test_aprobar_pone_asterisco_y_revertir_lo_quita(app, admin_client):
    ruta, operario, jefe = preparar(app, admin_client)
    original = ruta.read_bytes()
    respuesta = avisar(operario, [{'fila': 2, 'lado': 'de'}, {'fila': 4, 'lado': 'para'}])
    assert respuesta.status_code == 201
    identificador = respuesta.get_json()['id']

    # Pendiente: nada cambia, el terminal se sigue pidiendo
    assert ruta.read_bytes() == original
    assert ('W1', 'de') in cables_de_t1(admin_client)
    pendientes = admin_client.get('/api/engastado/propuestas/cables-pendientes?archivo=corte.xlsx&terminal=t1').get_json()
    assert sorted((c['fila'], c['lado']) for c in pendientes['cables']) == [(2, 'de'), (4, 'para')]
    assert admin_client.get('/api/engastado/propuestas/pendientes').get_json()['pendientes'] == 1

    respuesta = jefe.post(f'/api/engastado/propuestas/{identificador}/resolver',
                          json={'accion': 'aprobar', 'motivo': 'Comprobado en plano'})
    assert respuesta.status_code == 200, respuesta.get_json()
    hoja = load_workbook(ruta)['Format']
    assert hoja['D2'].value == 'X1*'
    assert hoja['F4'].value == 'Y3*'
    assert Path(str(ruta) + f'.backup_engastado_propuesta_{identificador}').read_bytes() == original
    restantes = cables_de_t1(admin_client)
    assert ('W1', 'de') not in restantes and ('W3', 'para') not in restantes
    assert ('W2', 'doble') in restantes

    respuesta = jefe.post(f'/api/engastado/propuestas/{identificador}/revertir', json={'motivo': 'Era un error'})
    assert respuesta.status_code == 200, respuesta.get_json()
    hoja = load_workbook(ruta)['Format']
    assert hoja['D2'].value == 'X1'
    assert hoja['F4'].value == 'Y3'
    assert ('W1', 'de') in cables_de_t1(admin_client)
    propuesta = jefe.get('/api/engastado/propuestas').get_json()['propuestas'][0]
    assert propuesta['estado'] == 'revertida'
    assert {c['activo'] for c in propuesta['cables']} == {0}


def test_revertir_no_toca_asteriscos_ajenos(app, admin_client):
    ruta, operario, jefe = preparar(app, admin_client)
    identificador = avisar(operario, [{'fila': 2, 'lado': 'de'}]).get_json()['id']
    jefe.post(f'/api/engastado/propuestas/{identificador}/resolver', json={'accion': 'aprobar', 'motivo': 'Comprobado'})
    libro = load_workbook(ruta)
    libro['Format']['I2'] = None   # alguien borra la marca: el '*' ya no es de la propuesta
    libro.save(ruta)
    assert jefe.post(f'/api/engastado/propuestas/{identificador}/revertir', json={'motivo': 'Revertir'}).status_code == 200
    assert load_workbook(ruta)['Format']['D2'].value == 'X1*'


def test_rechazar_no_cambia_el_excel_y_no_se_resuelve_dos_veces(app, admin_client):
    ruta, operario, jefe = preparar(app, admin_client)
    original = ruta.read_bytes()
    identificador = avisar(operario, [{'fila': 2, 'lado': 'de'}]).get_json()['id']
    assert operario.post(f'/api/engastado/propuestas/{identificador}/resolver',
                         json={'accion': 'aprobar', 'motivo': 'Sin admin'}).status_code == 401
    assert jefe.post(f'/api/engastado/propuestas/{identificador}/resolver',
                     json={'accion': 'rechazar', 'motivo': 'Si que va'}).status_code == 200
    assert ruta.read_bytes() == original
    assert jefe.post(f'/api/engastado/propuestas/{identificador}/resolver',
                     json={'accion': 'aprobar', 'motivo': 'Otra vez'}).status_code == 409
    assert jefe.post(f'/api/engastado/propuestas/{identificador}/revertir',
                     json={'motivo': 'No aprobada'}).status_code == 409


def test_aviso_duplicado_y_fila_que_no_cuadra(app, admin_client):
    ruta, operario, jefe = preparar(app, admin_client)
    assert avisar(operario, [{'fila': 2, 'lado': 'de'}]).status_code == 201
    assert avisar(operario, [{'fila': 2, 'lado': 'de'}]).status_code == 409
    # La fila 2 lleva T2 en Para, no T1
    assert avisar(operario, [{'fila': 2, 'lado': 'para'}]).status_code == 400
    assert avisar(operario, [{'fila': 99, 'lado': 'de'}]).status_code == 400
    assert avisar(operario, []).status_code == 400


def test_ed_nueva_hereda_el_asterisco_aunque_cambien_las_filas(app, admin_client):
    ruta, operario, jefe = preparar(app, admin_client)
    identificador = avisar(operario, [{'fila': 2, 'lado': 'de'}]).get_json()['id']
    jefe.post(f'/api/engastado/propuestas/{identificador}/resolver', json={'accion': 'aprobar', 'motivo': 'Comprobado'})

    nueva = excel(FILAS, extra_arriba=2)   # W1 pasa de la fila 2 a la 4
    respuesta = admin_client.post('/api/upload', data={'file': (io.BytesIO(nueva), 'corte.xlsx')},
                                  content_type='multipart/form-data')
    assert respuesta.status_code == 200, respuesta.get_json()
    assert respuesta.get_json()['engastado_heredado'] == {'heredados': 1, 'no_heredados': 0}
    hoja = load_workbook(ruta)['Format']
    assert hoja['D4'].value == 'X1*'
    assert hoja['D2'].value == 'Z9'
    assert ('W1', 'de') not in cables_de_t1(admin_client)
    cable = jefe.get('/api/engastado/propuestas').get_json()['propuestas'][0]['cables'][0]
    assert (cable['fila'], cable['activo'], cable['resultado']) == (4, 1, 'heredado')

    # Y revertir después de heredar quita el '*' de la fila nueva
    assert jefe.post(f'/api/engastado/propuestas/{identificador}/revertir', json={'motivo': 'Revertir'}).status_code == 200
    assert load_workbook(ruta)['Format']['D4'].value == 'X1'


def test_ed_nueva_sin_el_cable_no_lo_hereda(app, admin_client):
    ruta, operario, jefe = preparar(app, admin_client)
    identificador = avisar(operario, [{'fila': 2, 'lado': 'de'}]).get_json()['id']
    jefe.post(f'/api/engastado/propuestas/{identificador}/resolver', json={'accion': 'aprobar', 'motivo': 'Comprobado'})
    respuesta = admin_client.post('/api/upload', data={'file': (io.BytesIO(excel(FILAS[1:])), 'corte.xlsx')},
                                  content_type='multipart/form-data')
    assert respuesta.get_json()['engastado_heredado'] == {'heredados': 0, 'no_heredados': 1}
    cable = jefe.get('/api/engastado/propuestas').get_json()['propuestas'][0]['cables'][0]
    assert (cable['activo'], cable['resultado']) == (0, 'no_heredado_no_encontrado')


def test_filas_identicas_no_se_bloquean_a_ciegas():
    contenido = excel([FILAS[0], FILAS[0]])
    regla = {'lado': 'de', 'terminal': 'T1', 'cod_cable': 'C1', 'cable_marca': 'W1',
             'elemento': 'X1', 'serie': '', 'fila': None, 'propuesta_id': 7}
    nuevo, resultados = aplicar_bloqueos(contenido, 'corte.xlsx', [regla])
    assert resultados == [{'fila': None, 'resultado': 'ambiguo'}]
    hoja = load_workbook(io.BytesIO(nuevo))['Format']
    assert hoja['D2'].value == 'X1' and hoja['D3'].value == 'X1'
