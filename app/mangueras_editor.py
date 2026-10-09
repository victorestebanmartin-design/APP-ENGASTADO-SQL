import hashlib
import io
import os
import re
from copy import copy, deepcopy

from openpyxl import load_workbook

from app.excel_manager import (
    ExcelManager, _normalizar_texto_columna, _parse_instrucciones, _parse_pelado,
    _parse_retractiles, _tokens_invalidos_instrucciones, _serie_str,
)


PREPARACION_COLUMNAS = (
    'Instrucciones Mangueras DE', 'Instrucciones Mangueras PARA',
    'Retractil DE', 'Retractil PARA',
)


def listar_biblioteca_retractiles(carpeta):
    codigos = set()
    if not os.path.isdir(carpeta):
        return []
    manager = ExcelManager(carpeta)
    for nombre in os.listdir(carpeta):
        if not nombre.lower().endswith(('.xlsx', '.xlsm', '.xls')):
            continue
        try:
            for fila in manager.get_mangueras(nombre):
                for lado in ('retractil_de', 'retractil_para'):
                    codigos.update(retractil['codigo'] for retractil in fila[lado])
        except Exception:
            continue
    return sorted(codigos, key=str.casefold)


def _texto(valor):
    return '' if valor is None else str(valor).strip()


def _columnas(hoja):
    columnas = {}
    for celda in hoja[1]:
        if celda.value is not None:
            clave = _normalizar_texto_columna(celda.value)
            if clave in columnas:
                raise ValueError(f'Cabecera duplicada: {celda.value}')
            columnas[clave] = celda.column
    return columnas


def _valor(hoja, columnas, fila, *nombres):
    for nombre in nombres:
        indice = columnas.get(_normalizar_texto_columna(nombre))
        if indice:
            return _texto(hoja.cell(fila, indice).value)
    return ''


def _valor_texto(hoja, columnas, fila, nombre):
    indice = columnas.get(_normalizar_texto_columna(nombre))
    if not indice:
        return ''
    valor = hoja.cell(fila, indice).value
    return '' if valor is None else str(valor)


def _marca_base_asociada(marca):
    numero = re.fullmatch(r'(.+?)\(\s*([A-Za-z]+|\d+)\s*\)', marca, flags=re.I)
    if numero:
        base = numero.group(1).strip()
        base_sin_numero = re.sub(r'(?<=[A-Za-z])\d+$', '', base)
        return base_sin_numero.upper(), numero.group(2).upper()
    numero = re.fullmatch(r'(.+)-(\d+)', marca)
    if numero:
        return numero.group(1).strip().upper(), numero.group(2)
    return None, None


def _manguera_referenciada(observaciones):
    coincidencia = re.search(
        r'\bMANG(?:UERA)?\.?\s*(.+?)\s*\(\s*(\d+|S)\s*\)(?=\s|$)',
        observaciones or '', re.I,
    )
    if not coincidencia:
        return None, None
    return coincidencia.group(1).strip().upper(), coincidencia.group(2).upper()


def _clave_paquete_renfe(codigo, elemento, serie, seccion, observaciones):
    grupo = re.sub(r'\s+', ' ', (observaciones or '').strip()).casefold()
    if not grupo:
        return None
    return (codigo, elemento, serie, re.sub(r'\s+', '', seccion).upper(), grupo)


def _abrir(contenido, nombre):
    if not nombre.lower().endswith(('.xlsx', '.xlsm')):
        raise ValueError('Selecciona un Excel .xlsx o .xlsm. Convierte los .xls antes de abrirlos.')
    libro = load_workbook(io.BytesIO(contenido), keep_vba=nombre.lower().endswith('.xlsm'))
    hoja = libro['Format'] if 'Format' in libro.sheetnames else libro.worksheets[0]
    return libro, hoja, _columnas(hoja)


def _vinculos_mangueras(hoja, columnas):
    padres = {}
    candidatos = {}
    padres_paquete = {}
    activos = {}
    activos_marca = {}
    activos_paquete = {}
    for fila in range(2, hoja.max_row + 1):
        marca = (_valor(hoja, columnas, fila, 'Cable / Marca') or
                 _valor(hoja, columnas, fila, 'De Marca'))
        elemento = (_valor(hoja, columnas, fila, 'De Elemento Etiquetas') or
                _valor(hoja, columnas, fila, 'De Elemento')).rstrip('*').strip().upper()
        codigo = _valor(hoja, columnas, fila, 'Cod. cable').upper()
        serie = _serie_str(_valor(hoja, columnas, fila, 'Series'))
        try:
            longitud = float(_valor(hoja, columnas, fila, 'Longitud').replace(',', '.'))
        except ValueError:
            continue
        seccion = _valor(hoja, columnas, fila, 'Sección')
        if longitud > 0 and re.match(r'^\s*\d+\s*[xX\u00d7]', seccion) and marca and elemento and codigo:
            base = re.sub(r'-P$', '', marca, flags=re.I).upper()
            clave = (base, codigo, elemento, serie)
            padres[fila] = clave
            candidatos.setdefault(clave, []).append(fila)
            clave_paquete = _clave_paquete_renfe(
                codigo, elemento, serie, seccion, _valor(hoja, columnas, fila, 'Observaciones'))
            if clave_paquete:
                padres_paquete.setdefault(clave_paquete, []).append(fila)
        elif longitud == 0 and elemento and codigo:
            observaciones = _valor(hoja, columnas, fila, 'Observaciones')
            base, activo = _manguera_referenciada(observaciones)
            if base is not None:
                clave = (base, codigo, elemento, serie)
                activos.setdefault(clave, []).append((fila, activo))
            elif marca and _marca_base_asociada(marca)[0] is not None:
                base, activo = _marca_base_asociada(marca)
                clave = (base, codigo, elemento, serie)
                clave_paquete = _clave_paquete_renfe(
                    codigo, elemento, serie, seccion, observaciones)
                activos_marca.setdefault(clave, []).append((fila, activo, clave_paquete))
            else:
                numero = re.search(r'\(\s*(\d+)\s*\)', observaciones)
                if numero and marca:
                    base, activo = marca.upper(), numero.group(1)
                    clave = (base, codigo, elemento, serie)
                    activos.setdefault(clave, []).append((fila, activo))
                else:
                    numero = re.search(r'\(\s*(\d+|S)\s*\)\s*$', marca or '', re.I)
                    clave_paquete = _clave_paquete_renfe(
                        codigo, elemento, serie, seccion, observaciones)
                    if numero and clave_paquete:
                        activos_paquete.setdefault(clave_paquete, []).append(
                            (fila, numero.group(1).upper()))
    for clave, hijos in activos_marca.items():
        if candidatos.get(clave):
            activos.setdefault(clave, []).extend((fila, numero) for fila, numero, _ in hijos)
        else:
            for fila, numero, clave_paquete in hijos:
                if clave_paquete:
                    activos_paquete.setdefault(clave_paquete, []).append((fila, numero))
    resultado = {}
    for fila, clave in padres.items():
        encontrados = activos.get(clave, [])
        ambiguo = len(candidatos[clave]) != 1
        clave_paquete = _clave_paquete_renfe(
            clave[1], clave[2], clave[3],
            _valor(hoja, columnas, fila, 'Sección'),
            _valor(hoja, columnas, fila, 'Observaciones'))
        padres_grupo = padres_paquete.get(clave_paquete, []) if clave_paquete else []
        if not encontrados and padres_grupo:
            if len(padres_grupo) == 1:
                encontrados = activos_paquete.get(clave_paquete, [])
            elif activos_paquete.get(clave_paquete):
                ambiguo = True
        resultado[fila] = {'confirmados': bool(encontrados) and not ambiguo,
                           'ambiguo': ambiguo, 'activos': [], 'mallas': []}
        for fila_activo, numero in encontrados:
            grupo = 'mallas' if numero == 'S' else 'activos'
            resultado[fila][grupo].append({
                'fila': fila_activo, 'numero': numero,
                'cable_marca': _valor(hoja, columnas, fila_activo, 'Cable / Marca'),
                'de_elemento': _valor(hoja, columnas, fila_activo, 'De Elemento', 'De Elemento Etiquetas'),
                'para_elemento': _valor(hoja, columnas, fila_activo, 'Para Elemento'),
                'de_terminal': _valor(hoja, columnas, fila_activo, 'De Terminal'),
                'para_terminal': _valor(hoja, columnas, fila_activo, 'Para Terminal'),
                'bloqueo_automatico_de': _valor(hoja, columnas, fila_activo, 'Bloqueo Mangueras DE') in ('1', '1.0'),
                'bloqueo_automatico_para': _valor(hoja, columnas, fila_activo, 'Bloqueo Mangueras PARA') in ('1', '1.0'),
            })
    return resultado


def _columna_escritura(hoja, columnas, cabecera):
    clave = _normalizar_texto_columna(cabecera)
    if clave not in columnas:
        indice = hoja.max_column + 1
        nueva = hoja.cell(1, indice, cabecera)
        nueva._style = copy(hoja.cell(1, indice - 1)._style)
        hoja.column_dimensions[nueva.column_letter].width = 32
        columnas[clave] = indice
    hoja.cell(1, columnas[clave]).value = cabecera
    return columnas[clave]


def _columna_observaciones_mangueras(hoja, columnas):
    cabecera = 'Observaciones Mangueras'
    clave = _normalizar_texto_columna(cabecera)
    if clave in columnas:
        hoja.cell(1, columnas[clave]).value = cabecera
        return columnas[clave]
    indice = max(23, hoja.max_column + 1)
    nueva = hoja.cell(1, indice, cabecera)
    if indice > 1:
        nueva._style = copy(hoja.cell(1, indice - 1)._style)
    hoja.column_dimensions[nueva.column_letter].width = 42
    columnas[clave] = indice
    return indice


def _identidad_manguera_edicion(manguera):
    def normalizar(valor):
        texto = _texto(valor)
        return '' if texto.lower() in ('nan', 'none') else re.sub(r'\s+', ' ', texto).strip().casefold()

    identidad = tuple(normalizar(manguera.get(campo)) for campo in (
        'cable_marca', 'cod_cable', 'de_elemento', 'serie'))
    return identidad if all(identidad) else None


def _firma_lado(manguera, lado):
    instrucciones = manguera.get(lado)
    return (serializar_instrucciones(instrucciones) if instrucciones else '',
            _serializar_retractiles(manguera.get('retractil_' + lado) or []))


def _hay_preparacion(firma):
    return bool(firma[0] or firma[1])


def _anadir_aviso(manguera, aviso):
    actual = _texto(manguera.get('aviso_mangueras'))
    manguera['aviso_mangueras'] = f'{actual}\n{aviso}' if actual else aviso


def transferir_preparacion_edicion(contenido_origen, nombre_origen, contenido_nuevo, nombre_nuevo):
    """Hereda preparaciones al cambiar ED solo cuando marca/código/elemento/serie identifican una pareja única."""
    origen = leer_preparacion(contenido_origen, nombre_origen)['filas']
    destino = leer_preparacion(contenido_nuevo, nombre_nuevo)
    nuevos = destino['filas']
    origen_por_id = {}
    nuevos_por_id = {}
    for manguera in origen:
        identidad = _identidad_manguera_edicion(manguera)
        if identidad:
            origen_por_id.setdefault(identidad, []).append(manguera)
    for manguera in nuevos:
        identidad = _identidad_manguera_edicion(manguera)
        if identidad:
            nuevos_por_id.setdefault(identidad, []).append(manguera)

    cambios = []
    resumen = {'heredadas': 0, 'avisos': 0, 'ambiguas': 0}
    for identidad, anteriores in origen_por_id.items():
        anteriores_con_preparacion = [m for m in anteriores if any(
            _hay_preparacion(_firma_lado(m, lado)) for lado in ('de', 'para'))
            or m.get('observaciones_mangueras')]
        if not anteriores_con_preparacion:
            continue
        actuales = nuevos_por_id.get(identidad, [])
        if len(anteriores) != 1 or len(actuales) != 1:
            if actuales:
                aviso = (f'No se heredó la preparación de {nombre_origen}: '
                         'hay varias mangueras con la misma identidad. Revisar manualmente.')
                for actual in actuales:
                    _anadir_aviso(actual, aviso)
                    cambios.append(actual)
                resumen['ambiguas'] += 1
            continue

        anterior = anteriores_con_preparacion[0]
        actual = actuales[0]
        conflictos = []
        hubo_cambios = False
        for lado in ('de', 'para'):
            firma_anterior = _firma_lado(anterior, lado)
            firma_actual = _firma_lado(actual, lado)
            if not _hay_preparacion(firma_anterior):
                continue
            if not _hay_preparacion(firma_actual):
                actual[lado] = deepcopy(anterior[lado]) if anterior[lado] else None
                actual['retractil_' + lado] = deepcopy(anterior['retractil_' + lado])
                hubo_cambios = True
            elif firma_actual != firma_anterior:
                conflictos.append(f'Lado {lado.upper()}')

        observacion_anterior = anterior.get('observaciones_mangueras') or ''
        observacion_actual = actual.get('observaciones_mangueras') or ''
        if observacion_anterior and not observacion_actual:
            actual['observaciones_mangueras'] = observacion_anterior
            hubo_cambios = True
        elif observacion_anterior and observacion_actual != observacion_anterior:
            conflictos.append('observaciones')

        if conflictos:
            _anadir_aviso(actual,
                f'La ED nueva difiere de {nombre_origen} en {", ".join(conflictos)}. '
                'Se conservó la preparación nueva; revisar antes de aplicar.')
            resumen['avisos'] += 1
            hubo_cambios = True
        if hubo_cambios:
            cambios.append(actual)
            if not conflictos:
                resumen['heredadas'] += 1

    if not cambios:
        return io.BytesIO(contenido_nuevo), resumen
    resultado = exportar_preparacion(contenido_nuevo, nombre_nuevo, cambios, destino['revision'])
    return resultado, resumen


def _lados_preparados(hoja, columnas, fila):
    de = _valor(hoja, columnas, fila, PREPARACION_COLUMNAS[0])
    para = _valor(hoja, columnas, fila, PREPARACION_COLUMNAS[1])
    if de or para:
        instrucciones = {'de': _parse_instrucciones(de), 'para': _parse_instrucciones(para)}
    else:
        instrucciones = _parse_pelado(_valor(hoja, columnas, fila, 'Observaciones'))
    return {lado: (instrucciones[lado] or {}).get('pm') is not None for lado in ('de', 'para')}


def _actualizar_bloqueos(hoja, columnas, fila, vinculacion):
    if not vinculacion.get('confirmados'):
        return
    preparados = _lados_preparados(hoja, columnas, fila)
    for lado, cabecera_elemento, cabecera_terminal in (
            ('de', 'De Elemento', 'De Terminal'), ('para', 'Para Elemento', 'Para Terminal')):
        bloquear = not preparados[lado]
        marca_bloqueo = 'Bloqueo Mangueras ' + lado.upper()
        asociados = vinculacion['activos'] + vinculacion.get('mallas', [])
        for fila_destino in [fila] + [activo['fila'] for activo in asociados]:
            terminal = _valor(hoja, columnas, fila_destino, cabecera_terminal)
            if terminal.upper() in ('', 'S/T', 'NAN', 'NONE'):
                continue
            elemento = _valor(hoja, columnas, fila_destino, cabecera_elemento)
            automatico = _valor(hoja, columnas, fila_destino, marca_bloqueo) in ('1', '1.0')
            if bloquear and not elemento.endswith('*'):
                if lado == 'de' and not elemento:
                    elemento = _valor(hoja, columnas, fila_destino, 'De Elemento Etiquetas')
                hoja.cell(fila_destino, _columna_escritura(hoja, columnas, cabecera_elemento), elemento + '*')
                hoja.cell(fila_destino, _columna_escritura(hoja, columnas, marca_bloqueo), '1')
            elif not bloquear and automatico:
                hoja.cell(fila_destino, _columna_escritura(hoja, columnas, cabecera_elemento),
                          elemento[:-1] if elemento.endswith('*') else elemento)
                hoja.cell(fila_destino, _columna_escritura(hoja, columnas, marca_bloqueo)).value = None


def leer_preparacion(contenido, nombre):
    libro, hoja, columnas = _abrir(contenido, nombre)
    try:
        if _normalizar_texto_columna('Sección') not in columnas:
            raise ValueError('El Excel no contiene la columna Sección.')
        vinculos = _vinculos_mangueras(hoja, columnas)
        filas_activos = {activo['fila'] for vinculo in vinculos.values() if vinculo['confirmados']
                for activo in vinculo['activos'] + vinculo.get('mallas', [])}
        biblioteca_retractiles = set()
        for fila in range(2, hoja.max_row + 1):
            for cabecera in ('Retractil DE', 'Retráctil DE', 'Retractil PARA', 'Retráctil PARA'):
                valor = _valor(hoja, columnas, fila, cabecera)
                biblioteca_retractiles.update(
                    retractil['codigo'] for retractil in _parse_retractiles(valor)
                )
        filas = []
        for fila in range(2, hoja.max_row + 1):
            if fila in filas_activos:
                continue
            seccion = _valor(hoja, columnas, fila, 'Sección')
            inst_de = _valor(hoja, columnas, fila, PREPARACION_COLUMNAS[0])
            inst_para = _valor(hoja, columnas, fila, PREPARACION_COLUMNAS[1])
            ret_de = _valor(hoja, columnas, fila, PREPARACION_COLUMNAS[2])
            ret_para = _valor(hoja, columnas, fila, PREPARACION_COLUMNAS[3])
            obs = _valor(hoja, columnas, fila, 'Observaciones')
            sugerida = bool(re.match(r'^\s*\d+\s*[xX\u00d7]', seccion))
            if not sugerida and not any((inst_de, inst_para, ret_de, ret_para, '<-' in obs, '->' in obs)):
                continue
            if inst_de or inst_para:
                de = _parse_instrucciones(inst_de)
                para = _parse_instrucciones(inst_para)
            else:
                legacy = _parse_pelado(obs)
                de = legacy['de'] or _parse_instrucciones('')
                para = legacy['para'] or _parse_instrucciones('')
                partes = re.split(r'//|\$', obs)
                inst_de = next((parte.split('<-', 1)[1].strip() for parte in partes if '<-' in parte), '')
                inst_para = next((parte.split('->', 1)[0].strip() for parte in partes if '->' in parte and '<-' not in parte), '')
            for instrucciones, raw in ((de, inst_de), (para, inst_para)):
                instrucciones['otros_tokens'] = _tokens_invalidos_instrucciones(raw)
            filas.append({
                'fila': fila, 'sugerida': sugerida, 'seccion': seccion,
                'cable_marca': _valor(hoja, columnas, fila, 'Cable / Marca', 'De Marca'),
                'cod_cable': _valor(hoja, columnas, fila, 'Cod. cable'),
                'serie': _serie_str(_valor(hoja, columnas, fila, 'Series')),
                'de_elemento': _valor(hoja, columnas, fila, 'De Elemento Etiquetas', 'De Elemento'),
                'de_elemento_original': _valor(hoja, columnas, fila, 'De Elemento', 'De Elemento Etiquetas'),
                'para_elemento': _valor(hoja, columnas, fila, 'Para Elemento'),
                'de_terminal': _valor(hoja, columnas, fila, 'De Terminal'),
                'para_terminal': _valor(hoja, columnas, fila, 'Para Terminal'),
                'bloqueo_automatico_de': _valor(hoja, columnas, fila, 'Bloqueo Mangueras DE') in ('1', '1.0'),
                'bloqueo_automatico_para': _valor(hoja, columnas, fila, 'Bloqueo Mangueras PARA') in ('1', '1.0'),
                'de': de, 'para': para,
                'retractil_de': _parse_retractiles(ret_de),
                'retractil_para': _parse_retractiles(ret_para),
                'retractil_de_raw': ret_de, 'retractil_para_raw': ret_para,
                'observaciones_mangueras': _valor_texto(hoja, columnas, fila, 'Observaciones Mangueras'),
                'aviso_mangueras': _valor_texto(hoja, columnas, fila, 'Aviso Mangueras'),
                'vinculacion': vinculos.get(fila, {'confirmados': False, 'ambiguo': False, 'activos': []}),
                'campos': {str(hoja.cell(1, indice).value): _texto(hoja.cell(fila, indice).value)
                           for indice in columnas.values()},
            })
        return {'nombre': nombre, 'hoja': hoja.title, 'filas': filas,
            'biblioteca_retractiles': sorted(biblioteca_retractiles, key=str.casefold),
                'revision': hashlib.sha256(contenido).hexdigest()}
    finally:
        libro.close()


def _medida(valor):
    if valor is None or valor == '':
        return None
    if isinstance(valor, bool) or not re.fullmatch(r'\d{1,6}', str(valor)):
        raise ValueError('Las medidas deben ser enteros entre 0 y 999999 mm.')
    return str(int(valor))


def serializar_instrucciones(instrucciones):
    if not isinstance(instrucciones, dict):
        raise ValueError('Instrucciones no válidas.')
    tokens = []
    for campo, prefijo in (('pm', 'PM'), ('a_todos', 'A')):
        medida = _medida(instrucciones.get(campo))
        if medida is not None:
            tokens.append(prefijo + medida)
    modos = [campo for campo in ('m_cortar', 'm_mrs', 'm_mrc') if instrucciones.get(campo)]
    if len(modos) > 1 or (modos and instrucciones.get('m') not in (None, '')):
        raise ValueError('Selecciona un único tratamiento de malla.')
    if modos:
        modo = modos[0]
        if modo == 'm_cortar':
            tokens.append('M_CORTAR')
        else:
            tokens.append(('MRS' if modo == 'm_mrs' else 'MRC') +
                          (_medida(instrucciones.get(modo + '_medida')) or ''))
    else:
        medida = _medida(instrucciones.get('m'))
        if medida is not None:
            tokens.append('M' + medida)
    activos = instrucciones.get('a_especificos', {})
    if not isinstance(activos, dict):
        raise ValueError('Activos individuales no válidos.')
    for activo, valor in activos.items():
        if not re.fullmatch(r'[1-9]\d{0,3}', str(activo)):
            raise ValueError('El número de activo debe estar entre 1 y 9999.')
        medida = _medida(valor)
        if medida is None:
            raise ValueError('Falta la medida del activo.')
        tokens.append(f'A{activo}_{medida}')
    otros = instrucciones.get('otros_tokens', [])
    if not isinstance(otros, list) or any(not isinstance(token, str) or
                                          re.search(r'[/\r\n]', token) for token in otros):
        raise ValueError('Tokens adicionales no válidos.')
    tokens.extend(otros)
    return '/'.join(tokens)


def _serializar_retractiles(retractiles):
    if not isinstance(retractiles, list):
        raise ValueError('Lista de retráctiles no válida.')
    tokens = []
    for retractil in retractiles:
        if not isinstance(retractil, dict):
            raise ValueError('Retráctil no válido.')
        codigo = _texto(retractil.get('codigo'))
        medida = _medida(retractil.get('medida'))
        if not codigo or re.search(r'[/\r\n]', codigo) or codigo.startswith(('=', '+', '-', '@')) or medida is None:
            raise ValueError('Cada retráctil necesita código y medida, sin / en el código.')
        tokens.append(f'{codigo}_{medida}')
    return '/'.join(tokens)


def exportar_preparacion(contenido, nombre, cambios, revision):
    if revision != hashlib.sha256(contenido).hexdigest():
        raise ValueError('El Excel ha cambiado. Vuelve a cargarlo antes de descargar.')
    if not isinstance(cambios, list):
        raise ValueError('Los cambios deben ser una lista.')
    disponibles = {fila['fila'] for fila in leer_preparacion(contenido, nombre)['filas']}
    libro, hoja, columnas = _abrir(contenido, nombre)
    try:
        vinculos = _vinculos_mangueras(hoja, columnas)
        vistos = set()
        for cambio in cambios:
            if not isinstance(cambio, dict):
                raise ValueError('Cambio no válido.')
            fila = cambio.get('fila')
            if type(fila) is not int or fila not in disponibles or fila in vistos:
                raise ValueError('Fila de manguera no válida o repetida.')
            vistos.add(fila)
            observaciones = cambio.get('observaciones_mangueras', '')
            if not isinstance(observaciones, str):
                raise ValueError('Las observaciones de mangueras deben ser texto.')
            aviso = cambio.get('aviso_mangueras', '')
            if not isinstance(aviso, str):
                raise ValueError('El aviso de mangueras debe ser texto.')
            valores = (
                serializar_instrucciones(cambio.get('de')),
                serializar_instrucciones(cambio.get('para')),
                _serializar_retractiles(cambio.get('retractil_de')),
                _serializar_retractiles(cambio.get('retractil_para')),
            )
            valores = list(valores)
            for posicion, lado in ((2, 'de'), (3, 'para')):
                raw = _valor(hoja, columnas, fila, PREPARACION_COLUMNAS[posicion])
                if cambio.get('retractil_' + lado) == _parse_retractiles(raw):
                    valores[posicion] = raw
            for cabecera, valor in zip(PREPARACION_COLUMNAS, valores):
                celda = hoja.cell(fila, _columna_escritura(hoja, columnas, cabecera))
                celda.value = valor or None
                celda.data_type = 's'
            celda_observaciones = hoja.cell(fila, _columna_observaciones_mangueras(hoja, columnas))
            celda_observaciones.value = observaciones or None
            celda_observaciones.data_type = 's'
            clave_aviso = _normalizar_texto_columna('Aviso Mangueras')
            if aviso or clave_aviso in columnas:
                celda_aviso = hoja.cell(fila, _columna_escritura(hoja, columnas, 'Aviso Mangueras'))
                celda_aviso.value = aviso or None
                celda_aviso.data_type = 's'
            if not any(valores[:2]):
                indice_obs = columnas.get(_normalizar_texto_columna('Observaciones'))
                if indice_obs:
                    celda = hoja.cell(fila, indice_obs)
                    obs = _texto(celda.value)
                    if '<-' in obs or '->' in obs:
                        celda.value = ' // '.join(parte.strip() for parte in re.split(r'//|\$', obs)
                                                 if '<-' not in parte and '->' not in parte) or None
        for fila, vinculacion in vinculos.items():
            _actualizar_bloqueos(hoja, columnas, fila, vinculacion)
        salida = io.BytesIO()
        libro.save(salida)
        salida.seek(0)
        return salida
    finally:
        libro.close()


def _identidad_edicion_manguera(manguera):
    def normalizar(valor):
        texto = _texto(valor)
        if texto.lower() in ('nan', 'none'):
            return ''
        return re.sub(r'\s+', ' ', texto).strip().casefold()

    identidad = tuple(normalizar(manguera.get(campo)) for campo in (
        'cable_marca', 'cod_cable', 'de_elemento', 'serie'))
    return identidad if all(identidad) else None


def _firma_preparacion_lado(manguera, lado):
    instrucciones = manguera.get(lado)
    return (serializar_instrucciones(instrucciones) if instrucciones else '',
            _serializar_retractiles(manguera.get('retractil_' + lado) or []))


def _preparacion_lado_presente(firma):
    return bool(firma[0] or firma[1])


def _agregar_aviso(manguera, texto):
    aviso = manguera.get('aviso_mangueras') or ''
    if texto not in aviso:
        manguera['aviso_mangueras'] = (aviso + '\n' + texto).strip()


def transferir_preparacion_edicion(contenido_anterior, nombre_anterior,
                                   contenido_nuevo, nombre_nuevo):
    """Hereda preparación entre ED solo si la identidad de manguera es única en ambos Excels."""
    anteriores = leer_preparacion(contenido_anterior, nombre_anterior)['filas']
    nuevo = leer_preparacion(contenido_nuevo, nombre_nuevo)
    actuales = nuevo['filas']
    anteriores_por_identidad = {}
    actuales_por_identidad = {}
    for manguera in anteriores:
        identidad = _identidad_edicion_manguera(manguera)
        if identidad:
            anteriores_por_identidad.setdefault(identidad, []).append(manguera)
    for manguera in actuales:
        identidad = _identidad_edicion_manguera(manguera)
        if identidad:
            actuales_por_identidad.setdefault(identidad, []).append(manguera)

    cambios = []
    resumen = {'heredadas': 0, 'avisos': 0, 'ambiguas': 0}
    for identidad, grupo_anterior in anteriores_por_identidad.items():
        con_preparacion = [fila for fila in grupo_anterior if any(
            _preparacion_lado_presente(_firma_preparacion_lado(fila, lado))
            for lado in ('de', 'para')) or fila.get('observaciones_mangueras')]
        if not con_preparacion:
            continue
        grupo_nuevo = actuales_por_identidad.get(identidad, [])
        if len(grupo_anterior) != 1 or len(grupo_nuevo) != 1:
            if grupo_nuevo:
                aviso = (f'No se heredó la preparación de {nombre_anterior}: '
                         'la identidad de esta manguera no es única. Revisar manualmente.')
                for fila in grupo_nuevo:
                    _agregar_aviso(fila, aviso)
                    cambios.append(fila)
                resumen['ambiguas'] += 1
            continue

        fila_anterior = con_preparacion[0]
        fila_nueva = grupo_nuevo[0]
        conflictos = []
        hubo_cambios = False
        for lado in ('de', 'para'):
            anterior = _firma_preparacion_lado(fila_anterior, lado)
            actual = _firma_preparacion_lado(fila_nueva, lado)
            if not _preparacion_lado_presente(anterior):
                continue
            if not _preparacion_lado_presente(actual):
                fila_nueva[lado] = deepcopy(fila_anterior[lado]) if fila_anterior[lado] else None
                fila_nueva['retractil_' + lado] = deepcopy(fila_anterior['retractil_' + lado])
                hubo_cambios = True
            elif anterior != actual:
                conflictos.append('Lado ' + lado.upper())

        observaciones_anteriores = fila_anterior.get('observaciones_mangueras') or ''
        observaciones_nuevas = fila_nueva.get('observaciones_mangueras') or ''
        if observaciones_anteriores and not observaciones_nuevas:
            fila_nueva['observaciones_mangueras'] = observaciones_anteriores
            hubo_cambios = True
        elif observaciones_anteriores and observaciones_nuevas != observaciones_anteriores:
            conflictos.append('Observaciones')

        aviso_anterior = fila_anterior.get('aviso_mangueras') or ''
        if aviso_anterior and not fila_nueva.get('aviso_mangueras'):
            fila_nueva['aviso_mangueras'] = aviso_anterior
            hubo_cambios = True
        if conflictos:
            _agregar_aviso(
                fila_nueva,
                f'La ED nueva difiere de {nombre_anterior} en {", ".join(conflictos)}. '
                'Se conservaron los valores nuevos; revisar antes de aplicar.',
            )
            resumen['avisos'] += 1
            hubo_cambios = True
        if hubo_cambios:
            cambios.append(fila_nueva)
            if not conflictos:
                resumen['heredadas'] += 1

    if not cambios:
        return io.BytesIO(contenido_nuevo), resumen
    salida = exportar_preparacion(contenido_nuevo, nombre_nuevo, cambios, nuevo['revision'])
    return salida, resumen